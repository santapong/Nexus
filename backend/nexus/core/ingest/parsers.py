"""Document parsers — bytes in, normalized markdown + tables out (ADR-090).

Each concrete parser is synchronous CPU/IO work executed via
asyncio.to_thread so the event loop never blocks. Parsers never raise to
callers: failures are reported in ParsedDocument.status / .error so an
unparseable upload degrades to metadata-only instead of a 500.
"""

from __future__ import annotations

import asyncio
import io
from pathlib import PurePosixPath
from typing import Any

import structlog
from pydantic import BaseModel, Field

logger = structlog.get_logger()

# Cap extracted table rows so a million-row parquet doesn't explode context.
_MAX_TABLE_ROWS = 200
# Cap total extracted characters per document at parse time; the per-task
# context budget (settings.attachment_context_char_budget) trims further.
_MAX_TEXT_CHARS = 200_000


class ParsedDocument(BaseModel):
    """Normalized parse result cached on the attachments row."""

    markdown: str = ""
    tables: list[dict[str, Any]] = Field(default_factory=list)
    status: str = "parsed"  # parsed | failed | unsupported
    error: str | None = None


_EXTENSION_MIME: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".parquet": "application/vnd.apache.parquet",
    ".txt": "text/plain",
    ".md": "text/markdown",
}


def _kind_for(filename: str, mime_type: str) -> str:
    """Resolve the parser kind from extension first, then MIME type."""
    suffix = PurePosixPath(filename.lower()).suffix
    if suffix in {".pdf"}:
        return "pdf"
    if suffix in {".docx"}:
        return "docx"
    if suffix in {".xlsx", ".xlsm"}:
        return "xlsx"
    if suffix in {".csv"}:
        return "csv"
    if suffix in {".parquet"}:
        return "parquet"
    if suffix in {".txt", ".md", ".markdown", ".rst", ".log", ".json", ".yaml", ".yml"}:
        return "text"
    if mime_type == "application/pdf":
        return "pdf"
    if mime_type.startswith("text/"):
        return "text"
    return "unsupported"


def _rows_to_markdown(header: list[str], rows: list[list[Any]]) -> str:
    """Render rows as a markdown pipe table, capped at _MAX_TABLE_ROWS."""
    shown = rows[:_MAX_TABLE_ROWS]
    lines = [
        "| " + " | ".join(str(h) for h in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    lines.extend("| " + " | ".join(str(c) for c in row) + " |" for row in shown)
    if len(rows) > _MAX_TABLE_ROWS:
        lines.append(f"\n*…{len(rows) - _MAX_TABLE_ROWS} more rows omitted*")
    return "\n".join(lines)


def _parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"## Page {i + 1}\n\n{text.strip()}")
    markdown = "\n\n".join(pages)
    if not markdown.strip():
        return ParsedDocument(
            markdown="",
            status="parsed",
            error="No extractable text (scanned PDF? OCR is out of MVP scope)",
        )
    return ParsedDocument(markdown=markdown[:_MAX_TEXT_CHARS])


def _parse_docx(data: bytes) -> ParsedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        if style.startswith("Heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            parts.append(f"{'#' * min(int(level), 6)} {text}")
        else:
            parts.append(text)
    tables: list[dict[str, Any]] = []
    for t_idx, table in enumerate(document.tables):
        grid = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not grid:
            continue
        header, *rows = grid
        parts.append(_rows_to_markdown(header, rows))
        tables.append({"index": t_idx, "header": header, "row_count": len(rows)})
    return ParsedDocument(markdown="\n\n".join(parts)[:_MAX_TEXT_CHARS], tables=tables)


def _parse_xlsx(data: bytes) -> ParsedDocument:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    tables: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        rows = [
            ["" if c is None else c for c in row]
            for row in sheet.iter_rows(values_only=True, max_row=_MAX_TABLE_ROWS + 1)
        ]
        if not rows:
            continue
        header = [str(h) for h in rows[0]]
        body = [list(r) for r in rows[1:]]
        parts.append(f"## Sheet: {sheet.title}\n\n{_rows_to_markdown(header, body)}")
        tables.append({"sheet": sheet.title, "header": header, "row_count": len(body)})
    workbook.close()
    return ParsedDocument(markdown="\n\n".join(parts)[:_MAX_TEXT_CHARS], tables=tables)


def _parse_dataframe(df: Any, source: str) -> ParsedDocument:
    header = [str(c) for c in df.columns]
    rows = df.head(_MAX_TABLE_ROWS).astype(str).values.tolist()
    summary = f"## {source} — {len(df)} rows x {len(header)} columns\n\n"
    markdown = summary + _rows_to_markdown(header, rows)
    if len(df) > _MAX_TABLE_ROWS:
        markdown += f"\n\n*…{len(df) - _MAX_TABLE_ROWS} more rows omitted*"
    tables = [{"source": source, "header": header, "row_count": len(df)}]
    return ParsedDocument(markdown=markdown[:_MAX_TEXT_CHARS], tables=tables)


def _parse_csv(data: bytes) -> ParsedDocument:
    import pandas as pd

    df = pd.read_csv(io.BytesIO(data))
    return _parse_dataframe(df, "CSV")


def _parse_parquet(data: bytes) -> ParsedDocument:
    import pandas as pd

    df = pd.read_parquet(io.BytesIO(data))
    return _parse_dataframe(df, "Parquet")


def _parse_text(data: bytes) -> ParsedDocument:
    text = data.decode("utf-8", errors="replace")
    return ParsedDocument(markdown=text[:_MAX_TEXT_CHARS])


_PARSERS = {
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "xlsx": _parse_xlsx,
    "csv": _parse_csv,
    "parquet": _parse_parquet,
    "text": _parse_text,
}


def _parse_sync(data: bytes, filename: str, mime_type: str) -> ParsedDocument:
    kind = _kind_for(filename, mime_type)
    if kind == "unsupported":
        return ParsedDocument(
            status="unsupported",
            error=f"Unsupported document type: {filename} ({mime_type})",
        )
    try:
        return _PARSERS[kind](data)
    except Exception as exc:
        logger.warning(
            "document_parse_failed",
            filename=filename,
            mime_type=mime_type,
            kind=kind,
            error=str(exc),
        )
        return ParsedDocument(status="failed", error=str(exc))


async def parse_document(data: bytes, *, filename: str, mime_type: str) -> ParsedDocument:
    """Parse an uploaded document into normalized markdown + tables.

    Never raises — unparseable input returns status='failed'/'unsupported'
    with the reason in .error.

    Args:
        data: Raw file bytes.
        filename: Original filename (extension drives parser dispatch).
        mime_type: Client-declared MIME type (fallback dispatch signal).

    Returns:
        ParsedDocument with markdown, extracted tables, and parse status.
    """
    return await asyncio.to_thread(_parse_sync, data, filename, mime_type)
