"""Unit tests for document ingestion parsers (ADR-090).

Fixture files are authored in-test with the same libraries the parsers use
(python-docx, openpyxl, pandas) so no binary fixtures live in the repo.
"""

from __future__ import annotations

import io

import pytest

from nexus.core.ingest import parse_document


@pytest.mark.asyncio
async def test_docx_paragraphs_and_tables() -> None:
    import docx

    doc = docx.Document()
    doc.add_heading("Quarterly Report", level=1)
    doc.add_paragraph("Revenue grew 42 percent.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Revenue"
    table.rows[1].cells[1].text = "1.2M"
    buf = io.BytesIO()
    doc.save(buf)

    parsed = await parse_document(buf.getvalue(), filename="report.docx", mime_type="")

    assert parsed.status == "parsed"
    assert "# Quarterly Report" in parsed.markdown
    assert "Revenue grew 42 percent." in parsed.markdown
    assert "| Metric | Value |" in parsed.markdown
    assert parsed.tables and parsed.tables[0]["row_count"] == 1


@pytest.mark.asyncio
async def test_xlsx_sheets_to_markdown() -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Cost"])
    ws.append(["GPU", 2000])
    buf = io.BytesIO()
    wb.save(buf)

    parsed = await parse_document(buf.getvalue(), filename="budget.xlsx", mime_type="")

    assert parsed.status == "parsed"
    assert "## Sheet: Budget" in parsed.markdown
    assert "| Item | Cost |" in parsed.markdown
    assert "| GPU | 2000 |" in parsed.markdown


@pytest.mark.asyncio
async def test_csv_rows_and_row_cap() -> None:
    header = "name,score\n"
    rows = "".join(f"row{i},{i}\n" for i in range(250))

    parsed = await parse_document(
        (header + rows).encode(), filename="scores.csv", mime_type="text/csv"
    )

    assert parsed.status == "parsed"
    assert "| name | score |" in parsed.markdown
    assert "more rows omitted" in parsed.markdown  # capped at 200
    assert parsed.tables[0]["row_count"] == 250


@pytest.mark.asyncio
async def test_parquet_roundtrip() -> None:
    import pandas as pd

    df = pd.DataFrame({"city": ["Bangkok", "Osaka"], "pop_m": [11.2, 2.7]})
    buf = io.BytesIO()
    df.to_parquet(buf)

    parsed = await parse_document(buf.getvalue(), filename="cities.parquet", mime_type="")

    assert parsed.status == "parsed"
    assert "Bangkok" in parsed.markdown
    assert parsed.tables[0]["row_count"] == 2


@pytest.mark.asyncio
async def test_pdf_without_text_reports_note() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)

    parsed = await parse_document(buf.getvalue(), filename="blank.pdf", mime_type="")

    assert parsed.status == "parsed"
    assert parsed.markdown == ""
    assert parsed.error is not None and "No extractable text" in parsed.error


@pytest.mark.asyncio
async def test_plain_text_passthrough() -> None:
    parsed = await parse_document(b"hello notes", filename="notes.txt", mime_type="text/plain")
    assert parsed.status == "parsed"
    assert parsed.markdown == "hello notes"


@pytest.mark.asyncio
async def test_unsupported_type() -> None:
    parsed = await parse_document(
        b"PK\x03\x04", filename="archive.zip", mime_type="application/zip"
    )
    assert parsed.status == "unsupported"
    assert parsed.error is not None


@pytest.mark.asyncio
async def test_corrupt_file_fails_without_raising() -> None:
    parsed = await parse_document(b"not really an xlsx", filename="broken.xlsx", mime_type="")
    assert parsed.status == "failed"
    assert parsed.error
